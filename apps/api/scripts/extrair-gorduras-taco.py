# -*- coding: utf-8 -*-
"""
Extrai os ácidos graxos da planilha OFICIAL da TACO (NEPA/UNICAMP).

    python -I scripts/extrair-gorduras-taco.py Taco-4a-Edicao.xlsx > gorduras-taco.json

Por que Python e não Node: a casa já extrai TACO em Python
(`extrair-taco-pdf.py`), e ler xlsx em Node exigiria uma dependência nova no
pacote da API — que entraria no lockfile e na imagem do Docker por causa de um
script que roda uma vez.

------------------------------------------------------------------ A PLANILHA

Aba `AGtaco3`, cabeçalho em três linhas. As colunas que interessam:

    0  Número do Alimento      2  Saturados (g)
    3  Monoinsaturados (g)     4  Poliinsaturados (g)
   23  18:1t (g)              24  18:2t (g)        <- os trans

------------------------------------------------------- AS CONVENÇÕES DA TACO

    Tr  = traço, abaixo do limite de quantificação  -> 0
    NA  = não analisado                             -> null, NUNCA 0
    ""  = idem                                      -> null

A distinção entre `Tr` e `NA` é a razão de ser deste script. Gravar NA como
zero afirmaria que o alimento não tem gordura saturada — e um plano de queijo
somando "0 g de saturada" é exatamente o que não pode acontecer.

Mono e poli saem no JSON mesmo sem coluna no banco ainda: o dado é o mesmo
trabalho de extrair, e deixá-lo de fora obrigaria a repetir tudo depois.
"""
import io
import json
import sys

import openpyxl

sys.stdout.reconfigure(encoding="utf-8")

COL = {
    "numero": 0,
    "saturados": 2,
    "monoinsaturados": 3,
    "poliinsaturados": 4,
    "trans_18_1t": 23,
    "trans_18_2t": 24,
}


def valor(c):
    """Número, 0 para traço, None para não analisado."""
    if c is None:
        return None
    if isinstance(c, (int, float)):
        return float(c)
    t = str(c).strip()
    if not t:
        return None
    if t.upper().startswith("TR"):
        return 0.0
    if t.upper().startswith("NA"):
        return None
    try:
        return float(t.replace(",", "."))
    except ValueError:
        return None


def main(caminho):
    wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    if "AGtaco3" not in wb.sheetnames:
        print(json.dumps({"erro": f"sem aba AGtaco3; ha {wb.sheetnames}"}))
        return 1
    ws = wb["AGtaco3"]

    saida = {}
    for linha in ws.iter_rows(min_row=4, values_only=True):
        numero = valor(linha[COL["numero"]] if len(linha) > COL["numero"] else None)
        # Linha de título de grupo ("Cereais e derivados") não tem número.
        if numero is None:
            continue

        def ler(chave):
            i = COL[chave]
            return valor(linha[i]) if len(linha) > i else None

        sat = ler("saturados")
        mono = ler("monoinsaturados")
        poli = ler("poliinsaturados")
        t1 = ler("trans_18_1t")
        t2 = ler("trans_18_2t")
        # Os dois NA: o total é desconhecido, não zero. Um só NA: soma o outro.
        trans = None if (t1 is None and t2 is None) else (t1 or 0.0) + (t2 or 0.0)

        if sat is None and mono is None and poli is None and trans is None:
            continue

        saida[str(int(numero))] = {
            "saturadaG": sat,
            "monoinsaturadaG": mono,
            "poliinsaturadaG": poli,
            "transG": trans,
        }

    wb.close()
    print(json.dumps(saida, ensure_ascii=False, indent=1))
    return 0


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("uso: python -I scripts/extrair-gorduras-taco.py <Taco-4a-Edicao.xlsx>")
        sys.exit(1)
    sys.exit(main(sys.argv[1]))
